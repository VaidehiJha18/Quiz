import io
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT

from backend.services.paper_generator_service import get_paper_full_details

# Helper to set thin table borders in Word
def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'<w:top w:val="{kwargs.get("top", "none")}" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'<w:left w:val="{kwargs.get("left", "none")}" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'<w:bottom w:val="{kwargs.get("bottom", "none")}" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'<w:right w:val="{kwargs.get("right", "none")}" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def export_paper_to_docx(paper_id):
    """
    Exports the generated paper as a formatted DOCX document matching the GSFC University layout.
    """
    paper = get_paper_full_details(paper_id)
    if not paper:
        raise Exception("Paper not found")

    doc = Document()

    # Page Margins: 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Enrollment Header
    p_enroll = doc.add_paragraph()
    p_enroll.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_enroll = p_enroll.add_run("Enrollment No: ____________________")
    run_enroll.font.name = 'Times New Roman'
    run_enroll.font.size = Pt(10)
    run_enroll.font.bold = True

    # 2. Institutional Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    
    r_uni = p_title.add_run("GSFC UNIVERSITY\n")
    r_uni.font.name = 'Times New Roman'
    r_uni.font.size = Pt(14)
    r_uni.font.bold = True

    r_school = p_title.add_run("SCHOOL OF TECHNOLOGY\n")
    r_school.font.name = 'Times New Roman'
    r_school.font.size = Pt(12)
    r_school.font.bold = True

    exam_label = "Mid Semester Examination" if paper['exam_type'] == 'mid_sem' else "End Semester Examination"
    r_exam = p_title.add_run(f"B.Tech. CSE, {exam_label}, Odd Session {paper.get('academic_year', '2025-26')}")
    r_exam.font.name = 'Times New Roman'
    r_exam.font.size = Pt(11)
    r_exam.font.bold = True

    # 3. Exam Metadata Table
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    meta_rows = [
        ("Semester: VII", f"Date: {paper['created_at'].strftime('%d/%m/%Y') if hasattr(paper.get('created_at'), 'strftime') else 'DD/MM/YYYY'}"),
        (f"Course Code: {paper.get('course_code', 'BTCS701')}", f"Time: {'10:00 AM to 11:00 AM' if paper['exam_type'] == 'mid_sem' else '10:00 AM to 12:00 PM'}"),
        (f"Course Name: {paper.get('course_name', paper.get('paper_title', ''))}", f"Total Marks: {paper['total_marks']}")
    ]

    for i, (col1, col2) in enumerate(meta_rows):
        row = table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width, c2.width = Inches(3.5), Inches(3.5)
        
        c1.text, c2.text = col1, col2
        for cell in (c1, c2):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(10)
                    run.font.bold = True
            set_cell_border(cell, top="single", bottom="single", left="single", right="single")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 4. Standard Instructions
    p_inst = doc.add_paragraph()
    r_inst_head = p_inst.add_run("Instructions:\n")
    r_inst_head.font.name = 'Times New Roman'
    r_inst_head.font.size = Pt(10)
    r_inst_head.font.bold = True
    
    instructions_text = (
        "1. Attempt all questions.\n"
        "2. Figures to the right indicate full marks.\n"
        "3. Make suitable assumptions wherever necessary."
    )
    r_inst_body = p_inst.add_run(instructions_text)
    r_inst_body.font.name = 'Times New Roman'
    r_inst_body.font.size = Pt(9.5)
    r_inst_body.font.italic = True
    p_inst.paragraph_format.space_after = Pt(10)

    # 5. Question Sections & Sub-Parts
    for sec in paper.get('sections', []):
        sec_table = doc.add_table(rows=1, cols=3)
        sec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Section Header Row
        hdr_cells = sec_table.rows[0].cells
        hdr_cells[0].width = Inches(0.8)
        hdr_cells[1].width = Inches(5.4)
        hdr_cells[2].width = Inches(0.8)

        p_qno = hdr_cells[0].paragraphs[0]
        r_qno = p_qno.add_run(sec['question_number'])
        r_qno.font.bold = True
        r_qno.font.size = Pt(11)

        p_ins = hdr_cells[1].paragraphs[0]
        r_ins = p_ins.add_run(sec.get('instructions', 'Answer the following:'))
        r_ins.font.italic = True
        r_ins.font.size = Pt(10)

        p_tot = hdr_cells[2].paragraphs[0]
        p_tot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_tot = p_tot.add_run(f"({str(sec['total_section_marks']).zfill(2)})")
        r_tot.font.bold = True
        r_tot.font.size = Pt(11)

        # Helper to append sub-question rows
        def add_sub_row(part):
            row_cells = sec_table.add_row().cells
            row_cells[0].width = Inches(0.8)
            row_cells[1].width = Inches(5.4)
            row_cells[2].width = Inches(0.8)

            p_label = row_cells[0].paragraphs[0]
            p_label.add_run(part['sub_label']).font.bold = True

            p_txt = row_cells[1].paragraphs[0]
            p_txt.add_run(part['question_txt'])

            if part.get('diagram_url') and os.path.exists(part['diagram_url']):
                p_txt.add_run("\n")
                doc.add_picture(part['diagram_url'], width=Inches(3.0))

            p_m = row_cells[2].paragraphs[0]
            p_m.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_m.add_run(f"({str(part['marks']).zfill(2)})")

        for part in sec.get('main_questions', []):
            add_sub_row(part)

        # Render OR Choice block if questions exist
        if sec.get('or_questions'):
            or_row = sec_table.add_row().cells
            p_or = or_row[1].paragraphs[0]
            p_or.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_or = p_or.add_run("OR")
            r_or.font.bold = True
            r_or.font.size = Pt(11)

            for part in sec['or_questions']:
                add_sub_row(part)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # End Paper Divider
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("*********")
    r_end.font.bold = True

    # Save to memory buffer
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream